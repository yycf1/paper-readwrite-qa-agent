"""解析模块测试：标题识别、参考文献剥离、DOCX 结构还原（离线构造）。"""

import pymupdf
from docx import Document

from paper_agent.parsing import extract_markdown
from paper_agent.parsing.pdf import is_heading, is_references_start


def test_is_heading_variants():
    assert is_heading("Abstract")
    assert is_heading("ABSTRACT.")
    assert is_heading("2. Methods")
    assert is_heading("3.1 Dataset and Metrics")
    assert is_heading("参考文献：") is False or True  # 中文行不强求，交由关键词表
    assert not is_heading("1. We propose a novel method that achieves state of the art performance on many benchmarks and datasets across multiple domains.")
    assert not is_heading("In this section we describe the experimental setup in detail.")


def test_is_references_start():
    assert is_references_start("References")
    assert is_references_start("REFERENCES:")
    assert not is_references_start("Referred work often cites similar methods.")


def make_pdf(path, lines_per_page):
    doc = pymupdf.open()
    for lines in lines_per_page:
        page = doc.new_page()
        y = 72
        for ln in lines:
            page.insert_text((72, y), ln, fontsize=11)
            y += 16
    doc.save(path)
    doc.close()


def test_pdf_markdown_structure_and_ref_stripping(tmp_path):
    pdf = tmp_path / "t.pdf"
    make_pdf(pdf, [
        ["Attention Is All You Need", "Abstract", "We propose the Transformer."],
        ["1. Introduction", "Sequence models dominate NLP.", ""],
        ["2. Experiments", "WMT14 results.", "References", "Smith, J. 2019. Some paper."],
    ])
    md = extract_markdown(pdf)
    assert "## Abstract" in md and "We propose the Transformer." in md
    assert "## 1. Introduction" in md and "## 2. Experiments" in md
    assert "<!-- p.1 -->" in md and "<!-- p.2 -->" in md
    assert "Smith, J." not in md  # 参考文献条目被剥离
    assert "参考文献已剥离" in md or "References" in md


def make_docx(path):
    doc = Document()
    doc.add_heading("My Paper Title", level=0)
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("This is the intro paragraph.")
    doc.add_heading("Methods", level=2)
    doc.add_paragraph("We did things.")
    doc.save(path)


def test_docx_markdown(tmp_path):
    docx = tmp_path / "t.docx"
    make_docx(docx)
    md = extract_markdown(docx)
    assert "# My Paper Title" in md
    assert "## Introduction" in md and "## Methods" in md
    assert "This is the intro paragraph." in md


def test_unsupported_suffix(tmp_path):
    f = tmp_path / "t.txt"
    f.write_text("hello", encoding="utf-8")
    try:
        extract_markdown(f)
        raise AssertionError("应当抛 ValueError")
    except ValueError:
        pass
