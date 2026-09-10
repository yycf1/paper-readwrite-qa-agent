"""DOCX → Markdown：利用 Word 内建标题样式（Heading 1/2/…）还原章节结构。"""

from __future__ import annotations

from docx import Document


def extract_markdown(docx_path) -> str:
    doc = Document(str(docx_path))
    lines: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style == "title":
            lines.append("# " + text)  # add_heading(level=0) 产出 Title 样式
        elif style.startswith("heading"):
            digits = "".join(ch for ch in style if ch.isdigit())
            level = min(int(digits) + 1, 6) if digits else 2
            lines.append("#" * level + " " + text)
        else:
            lines.append(text)
    return "\n\n".join(lines).strip() + "\n"
